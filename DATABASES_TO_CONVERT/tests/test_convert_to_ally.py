# DATABASES_TO_CONVERT/tests/test_convert_to_ally.py
import sys, tempfile, unittest
from pathlib import Path

import openpyxl as xl  # verifica dipendenza

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import convert_to_ally as conv


def make_wb(sheets):
    """sheets: {nome: (headers, [tuple])} -> Workbook apribile/salvabile."""
    wb = xl.Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
    return wb


class RoundtripTest(unittest.TestCase):
    def test_main_copia_tutti_i_fogli_senza_modifiche(self):
        src_wb = make_wb({
            "ALUNNI": (("EMAIL", "ATTIVO"), (("a@x.it", True),)),
            "DOCENTI": (("EMAIL",), (("d@x.it",),)),
        })
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            out_dir = Path(tmp) / "out"
            src_wb.save(src)
                # main deve creare out_dir, salvare l'xlsx col nome del sorgente
            written = conv.main(src, out_dir)
            self.assertEqual(written.parent, out_dir)
            self.assertEqual(written.name, src.name)
            got = xl.load_workbook(written)
            self.assertEqual(got.sheetnames, ["ALUNNI", "DOCENTI"])
            self.assertEqual([c.value for c in got["ALUNNI"][1]], ["EMAIL", "ATTIVO"])
            self.assertEqual([c.value for c in got["ALUNNI"][2]], ["a@x.it", True])


if __name__ == "__main__":
    unittest.main()
