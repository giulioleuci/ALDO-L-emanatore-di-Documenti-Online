#!/usr/bin/env python3
"""Allinea 'AS 2526 per ALDO.xlsx' alle specifiche SGSA (vedi spec 2026-08-24).

Scrive <out_dir>/<stesso nome sorgente> + REPORT_ALLINEAMENTO.md.
Non modifica mai il sorgente.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl as xl

SOURCE_DEFAULT = Path(__file__).resolve().parent / "AS 2526 per ALDO.xlsx"
OUT_DIR_DEFAULT = Path(__file__).resolve().parent / "ALLY"
REPORT_NAME = "REPORT_ALLINEAMENTO.md"


class ChangeLog:
    """Registro di tutte le modifiche applicate (una voce per cella)."""

    def __init__(self) -> None:
        self.entries: list[dict] = []

    def add(self, rule: str, sheet: str, row_key: str, column: str, old, new) -> None:
        self.entries.append(
            {"rule": rule, "sheet": sheet, "row": row_key,
             "column": column, "old": "" if old is None else old,
             "new": "" if new is None else new}
        )

    def count(self, rule: str) -> int:
        return sum(1 for e in self.entries if e["rule"] == rule)


def load_workbook(path: Path) -> xl.Workbook:
    """Carica il workbook con tipi nativi (date/bool), formule risolte."""
    return xl.load_workbook(path, data_only=True)


def apply_rules(wb: xl.Workbook, log: ChangeLog) -> None:
    """Punto di estensione: le regole dei task seguenti si agganciano qui."""


def _col_index(headers: tuple, name: str) -> int:
    return list(headers).index(name)


def collect_residuals(wb: xl.Workbook) -> list[dict]:
    """Violazioni FK residue DOPO le regole (solo quelle attese vanno bene).

    I fogli assenti dal workbook sono ignorati (trattati come nessun dato).
    """
    def has(sheet_name: str) -> bool:
        return sheet_name in wb.sheetnames

    def ids(sheet_name: str) -> set:
        if not has(sheet_name):
            return set()
        ws = wb[sheet_name]
        return {r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}

    confdoc_ids = ids("CONF_DOC")
    confmail_ids = ids("CONF_MAIL")
    alunni = ids("ALUNNI")
    out: list[dict] = []

    def tally(sheet, field, values):
        from collections import Counter
        for value, count in Counter(v for v in values if v).items():
            out.append({"sheet": sheet, "field": field,
                        "value": str(value), "count": count})

    if has("CARENZE"):
        ws = wb["CARENZE"]
        i_st = _col_index(tuple(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))), "STUDENTE")
        tally("CARENZE", "STUDENTE",
              (r[i_st] for r in ws.iter_rows(min_row=2, values_only=True) if r[i_st] not in alunni))

    if has("LOG_DOCS"):
        ws = wb["LOG_DOCS"]
        hdr = tuple(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        i_ig = _col_index(hdr, "ID_GEN")
        empties, unknown = [], []
        for r in ws.iter_rows(min_row=2, values_only=True):
            ig = r[i_ig]
            if ig is None or str(ig).strip() == "":
                empties.append("<VUOTO>")
            elif ig not in confdoc_ids:
                unknown.append(ig)
        tally("LOG_DOCS", "ID_GEN", empties + unknown)

    if has("LOG_FOLD"):
        ws = wb["LOG_FOLD"]
        hdr = tuple(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        i_ti = _col_index(hdr, "TIPO")
        tally("LOG_FOLD", "TIPO",
              (r[i_ti] for r in ws.iter_rows(min_row=2, values_only=True) if r[i_ti] not in confdoc_ids))

    if has("LOG_MAIL"):
        ws = wb["LOG_MAIL"]
        hdr = tuple(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        i_pa = _col_index(hdr, "PATTERN")
        tally("LOG_MAIL", "PATTERN",
              (r[i_pa] for r in ws.iter_rows(min_row=2, values_only=True) if r[i_pa] not in confmail_ids))
    return out


# Coppie (sheet, field) i cui residui sono ACCETTATI dalla spec (§5/§6).
EXPECTED_RESIDUAL_KEYS = {
    ("CARENZE", "STUDENTE"),
    ("LOG_DOCS", "ID_GEN"),
    ("LOG_FOLD", "TIPO"),
    ("LOG_MAIL", "PATTERN"),
}


def write_report(log: ChangeLog, residuals: list[dict], path: Path) -> None:
    lines: list[str] = [
        "# Report allineamento AS 2526 per ALDO",
        "",
        "Generato da `convert_to_ally.py` — vedi spec "
        "`docs/superpowers/specs/2026-08-24-database-ally-alignment-design.md`.",
        "",
        "## Modifiche",
        "",
    ]
    by_rule: dict[str, list[dict]] = {}
    for e in log.entries:
        by_rule.setdefault(e["rule"], []).append(e)
    if not by_rule:
        lines.append("(nessuna modifica)")
    for rule in sorted(by_rule):
        entries = by_rule[rule]
        lines += [f"### {rule} — {len(entries)} modifiche", "",
                  "| Foglio | Riga | Colonna | Vecchio | Nuovo |",
                  "| --- | --- | --- | --- | --- |"]
        for e in entries[:100]:
            lines.append(f"| {e['sheet']} | {e['row']} | {e['column']} "
                         f"| `{e['old']}` | `{e['new']}` |")
        if len(entries) > 100:
            lines.append(f"| … | … | … | … ({len(entries) - 100} omesse) | ")
        lines.append("")

    lines += ["## Residui non risolti (attesi)", ""]
    unexpected = [r for r in residuals
                  if (r["sheet"], r["field"]) not in EXPECTED_RESIDUAL_KEYS]
    for r in sorted(residuals, key=lambda r: (r["sheet"], r["field"], r["value"])):
        lines.append(f"- `{r['sheet']}.{r['field']}` = `{r['value']}` — {r['count']} righe")
    if unexpected:
        lines += ["", "### RESIDUI NON ATTESI (richiedono attenzione)", ""]
        for r in unexpected:
            lines.append(f"- `{r['sheet']}.{r['field']}` = `{r['value']}` — {r['count']} righe")

    lines += [
        "",
        "## Deviazioni accettate",
        "",
        "- `PIANI_PERSONALIZZATI.MODELLI`: 4 rami anziché 5 (ramo PEI non previsto per scelta "
        "dell'utente): una generazione PEI produce errore esplicito.",
        "- Audit trail storici rimappati ai nuovi ID su richiesta esplicita (violazione locale "
        "dell'immutabilità append-only).",
        "- Dati di test mantenuti: harry.potter/HOGWARTS, classe ALTRO ISTITUTO, righe "
        "TEST_DOC_TABELLE_*, righe fantasma vuote.",
        "- Timestamp storici in formato `dd/MM/yyyy HH:mm:ss` (le nuove scritture saranno ISO).",
        "- Chiavi PARAMETRI inerti presenti (`colonneMaterie`, `schede`, `anniCorso`, "
        "`configurazionePipeline`): non lette da alcuno strategy.",
        "- `DIPART.CAPO` lasciato vuoto in attesa di nomi reali.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(source: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(source)
    log = ChangeLog()
    apply_rules(wb, log)
    dest = out_dir / source.name
    wb.save(dest)
    write_report(log, collect_residuals(load_workbook(dest)), out_dir / REPORT_NAME)
    return dest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE_DEFAULT)
    parser.add_argument("--out-dir", type=Path, default=OUT_DIR_DEFAULT)
    args = parser.parse_args()
    print(f"Scritto: {main(args.source, args.out_dir)}")
